from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from ragagent.chunking import chunk_document
from ragagent.documents import DocumentError, parse_document, safe_filename


def test_markdown_parser_preserves_sections_and_offsets(tmp_path: Path) -> None:
    path = tmp_path / "guide.md"
    path.write_text("# 安装\n先安装依赖。\n\n# 使用\n运行服务并上传文档。", encoding="utf-8")

    parsed = parse_document("doc_test", "guide.md", path)

    assert [section.section for section in parsed.sections] == ["安装", "使用"]
    assert parsed.sections[0].text == "先安装依赖。"
    chunks = chunk_document(parsed, target_chars=10, overlap_chars=2)
    assert chunks
    assert chunks[0].source.document_id == "doc_test"
    assert chunks[0].contextual_text.startswith("Document: guide.md")


def test_safe_filename_removes_path_and_dangerous_characters() -> None:
    assert safe_filename("../../weird:name?.txt") == "weird_name_.txt"


def test_empty_text_document_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("   ", encoding="utf-8")
    with pytest.raises(DocumentError, match="No extractable text"):
        parse_document("doc_empty", "empty.txt", path)


def test_utf8_text_parser_preserves_content(tmp_path: Path) -> None:
    path = tmp_path / "policy.txt"
    path.write_text("关键安全事件必须在四小时内上报。", encoding="utf-8")
    parsed = parse_document("doc_txt", path.name, path)
    assert parsed.sections[0].section == "Document"
    assert "四小时" in parsed.text


def test_docx_parser_preserves_heading(tmp_path: Path) -> None:
    path = tmp_path / "policy.docx"
    document = DocxDocument()
    document.add_heading("Incident response", level=1)
    document.add_paragraph("Critical incidents must be reported within four hours.")
    document.save(path)
    parsed = parse_document("doc_docx", path.name, path)
    assert parsed.sections[0].section == "Incident response"
    assert "four hours" in parsed.text


def test_pdf_parser_preserves_page_number(tmp_path: Path) -> None:
    path = tmp_path / "policy.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Annual leave is twelve days.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as handle:
        writer.write(handle)
    parsed = parse_document("doc_pdf", path.name, path)
    assert parsed.sections[0].page == 1
    assert "twelve days" in parsed.text


def test_corrupt_pdf_reports_parse_failure(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a pdf")
    with pytest.raises(DocumentError, match="PDF could not be opened"):
        parse_document("doc_broken", path.name, path)
