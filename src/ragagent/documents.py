from __future__ import annotations

import hashlib
import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}


class DocumentError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedSection:
    text: str
    page: int | None
    section: str | None
    start_char: int
    end_char: int


@dataclass(frozen=True)
class ParsedDocument:
    document_id: str
    filename: str
    sections: list[ParsedSection]

    @property
    def text(self) -> str:
        return "\n\n".join(section.text for section in self.sections)

    def to_json(self) -> str:
        return json.dumps(
            {
                "document_id": self.document_id,
                "filename": self.filename,
                "sections": [asdict(section) for section in self.sections],
            },
            ensure_ascii=False,
        )

    @classmethod
    def from_path(cls, path: Path) -> ParsedDocument:
        payload = json.loads(path.read_text(encoding="utf-8"))
        return cls(
            document_id=payload["document_id"],
            filename=payload["filename"],
            sections=[ParsedSection(**section) for section in payload["sections"]],
        )


def safe_filename(filename: str) -> str:
    name = Path(filename).name.replace("\x00", "").strip()
    name = re.sub(r"[\\/:*?\"<>|]", "_", name)
    if not name or name in {".", ".."}:
        raise DocumentError("The uploaded filename is invalid")
    return name[:240]


def content_sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def validate_extension(filename: str) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentError(f"Unsupported file type {extension!r}; supported: {supported}")
    return extension


def parse_document(document_id: str, filename: str, path: Path) -> ParsedDocument:
    extension = validate_extension(filename)
    if extension == ".pdf":
        sections = _parse_pdf(path)
    elif extension == ".docx":
        sections = _parse_docx(path)
    else:
        sections = _parse_text(path, markdown=extension in {".md", ".markdown"})
    sections = [section for section in sections if section.text.strip()]
    if not sections:
        raise DocumentError("No extractable text was found in the document")
    return ParsedDocument(document_id=document_id, filename=filename, sections=sections)


def _parse_pdf(path: Path) -> list[ParsedSection]:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # pypdf exposes several format-specific errors
        raise DocumentError(f"PDF could not be opened: {exc}") from exc
    sections: list[ParsedSection] = []
    cursor = 0
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as exc:
            raise DocumentError(f"PDF page {page_number} could not be extracted: {exc}") from exc
        if not text:
            continue
        sections.append(
            ParsedSection(
                text=text,
                page=page_number,
                section=f"Page {page_number}",
                start_char=cursor,
                end_char=cursor + len(text),
            )
        )
        cursor += len(text) + 2
    return sections


def _parse_docx(path: Path) -> list[ParsedSection]:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise DocumentError(f"DOCX could not be opened: {exc}") from exc
    sections: list[ParsedSection] = []
    current_heading = "Document"
    buffer: list[str] = []
    cursor = 0

    def flush() -> None:
        nonlocal cursor
        text = "\n".join(buffer).strip()
        if text:
            sections.append(
                ParsedSection(
                    text=text,
                    page=None,
                    section=current_heading,
                    start_char=cursor,
                    end_char=cursor + len(text),
                )
            )
            cursor += len(text) + 2
        buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name.lower() if paragraph.style else ""
        if style_name.startswith("heading"):
            flush()
            current_heading = text
        else:
            buffer.append(text)
    flush()
    return sections


def _parse_text(path: Path, *, markdown: bool) -> list[ParsedSection]:
    try:
        raw = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            raw = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentError("Text file must be UTF-8 encoded") from exc
    if not markdown:
        text = raw.strip()
        return [
            ParsedSection(
                text=text,
                page=None,
                section="Document",
                start_char=0,
                end_char=len(text),
            )
        ]

    sections: list[ParsedSection] = []
    heading = "Document"
    buffer: list[str] = []
    section_start = 0
    cursor = 0

    def flush() -> None:
        nonlocal section_start
        text = "\n".join(buffer).strip()
        if text:
            sections.append(
                ParsedSection(
                    text=text,
                    page=None,
                    section=heading,
                    start_char=section_start,
                    end_char=section_start + len(text),
                )
            )
        buffer.clear()
        section_start = cursor

    for line in raw.splitlines():
        match = re.match(r"^#{1,6}\s+(.+?)\s*$", line)
        if match:
            flush()
            heading = match.group(1)
            section_start = cursor + len(line) + 1
        else:
            buffer.append(line)
        cursor += len(line) + 1
    flush()
    return sections
